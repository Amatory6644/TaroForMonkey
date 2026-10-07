import io
from zipfile import ZipFile

from sqlalchemy import func, select

from telok.domain import DomainError
from telok.models import TaskArtifact


def relevant_memory(session, project_id, query):
    rank = func.ts_rank_cd(
        func.to_tsvector("simple", TaskArtifact.body.cast(__import__("sqlalchemy").Text)),
        func.plainto_tsquery("simple", query[:1000]),
    )
    rows = session.execute(
        select(TaskArtifact)
        .where(TaskArtifact.project_id == project_id, rank > 0)
        .order_by(rank.desc(), TaskArtifact.created_at.desc())
        .limit(3)
    ).scalars()
    return [{"stage": r.stage, "body": str(r.body)[:2500]} for r in rows]


def extract_document(data: bytes, filename: str):
    if len(data) > 10_000_000:
        raise DomainError("Документ превышает 10 MB.")
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 100:
            raise DomainError("PDF зашифрован или превышает 100 страниц.")
        value = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == "docx":
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30_000_000:
                raise DomainError("Слишком большой распакованный документ.")
        from docx import Document

        doc = Document(io.BytesIO(data))
        value = "\n".join(p.text for p in doc.paragraphs)
        value += "\n" + "\n".join(
            " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
        )
    elif ext in {"txt", "md", "csv"}:
        value = data.decode("utf-8-sig", errors="replace")
    else:
        raise DomainError("Поддерживаются PDF, DOCX, TXT, MD, CSV.")
    if not value.strip():
        raise DomainError("Текст не извлечён. Для скана нужен OCR.")
    return {"filename": filename[:150], "text": value[:50000], "truncated": len(value) > 50000}
